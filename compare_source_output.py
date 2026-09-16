import sys
from pathlib import Path
from docx import Document

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass


SOURCE_FILE = Path(
    r"D:\Bangla_Document_Converter\samples\Note_TEC.docx"
)

OUTPUT_FILE = Path(
    r"D:\Bangla_Document_Converter\output\Note_TEC_Unicode_Nikosh.docx"
)


def get_all_paragraphs(document):
    paragraphs = []

    # Body paragraphs
    for paragraph in document.paragraphs:
        paragraphs.append(paragraph.text)

    # Table paragraphs
    def collect_table(table):
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraphs.append(paragraph.text)

                for nested_table in cell.tables:
                    collect_table(nested_table)

    for table in document.tables:
        collect_table(table)

    return paragraphs


def main():

    print("=" * 80)
    print("SOURCE vs OUTPUT TEXT COMPARISON")
    print("=" * 80)

    print(f"Source : {SOURCE_FILE}")
    print(f"Output : {OUTPUT_FILE}")
    print()

    if not SOURCE_FILE.exists():
        raise FileNotFoundError(
            f"Source file not found:\n{SOURCE_FILE}"
        )

    if not OUTPUT_FILE.exists():
        raise FileNotFoundError(
            f"Output file not found:\n{OUTPUT_FILE}"
        )

    source_doc = Document(SOURCE_FILE)
    output_doc = Document(OUTPUT_FILE)

    source_paragraphs = get_all_paragraphs(source_doc)
    output_paragraphs = get_all_paragraphs(output_doc)

    print(f"Source text blocks : {len(source_paragraphs)}")
    print(f"Output text blocks : {len(output_paragraphs)}")
    print()

    print("=" * 80)
    print("NON-EMPTY TEXT COMPARISON")
    print("=" * 80)

    max_len = max(
        len(source_paragraphs),
        len(output_paragraphs),
    )

    differences = 0

    for i in range(max_len):

        source_text = (
            source_paragraphs[i]
            if i < len(source_paragraphs)
            else ""
        )

        output_text = (
            output_paragraphs[i]
            if i < len(output_paragraphs)
            else ""
        )

        if not source_text.strip() and not output_text.strip():
            continue

        print()
        print("-" * 80)
        print(f"TEXT BLOCK {i + 1}")
        print("-" * 80)

        print("SOURCE:")
        print(repr(source_text))

        print()

        print("OUTPUT:")
        print(repr(output_text))

        if source_text != output_text:
            differences += 1

    print()
    print("=" * 80)
    print("COMPARISON SUMMARY")
    print("=" * 80)

    print(
        f"Different text blocks : {differences}"
    )

    print()
    print(
        "This comparison is diagnostic only."
    )
    print(
        "No files were modified."
    )

    print("=" * 80)


if __name__ == "__main__":
    try:
        main()

    except Exception as exc:
        print()
        print("[ERROR]")
        print(str(exc))