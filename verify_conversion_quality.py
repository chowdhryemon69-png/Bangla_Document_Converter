import sys
from pathlib import Path
from docx import Document

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

OUTPUT_FILE = Path(r"D:\Bangla_Document_Converter\output\Note_TEC_Unicode_Nikosh.docx")

REQUIRED_PHRASES = [
    "টোকাসমূহের",
    "সেই মোতাবেক",
    "খোলা",
    "মোট",
    "মর্মে",
    "পোষণ",
    "প্রেরণ",
    "রেস্পন্সিভনেস",
    "মোঃ",
]

CORRUPTED_ARTEFACTS = [
    "েটাকাসমূেহর",
    "েসই েমাতাবেক",
    "েখালা",
    "েমাট",
    "মমের্",
    "েপাষণ",
    "েপ্ররণ",
    "েরস্পন্সিভনেস",
    "েমাঃ",
]

def main():
    print("=" * 80)
    print("CONVERSION QUALITY & ARTEFACT VERIFICATION TEST")
    print("=" * 80)
    print(f"Inspecting file: {OUTPUT_FILE}")
    print()

    if not OUTPUT_FILE.exists():
        print(f"[FAIL] Output file does not exist: {OUTPUT_FILE}")
        sys.exit(1)

    doc = Document(OUTPUT_FILE)

    # Collect all text
    paragraphs_text = []
    for p in doc.paragraphs:
        paragraphs_text.append(p.text)
    for t in doc.tables:
        def extract_table(tbl):
            for row in tbl.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        paragraphs_text.append(p.text)
                    for nested in cell.tables:
                        extract_table(nested)
        extract_table(t)

    full_text = "\n".join(paragraphs_text)

    all_passed = True

    print("--- 1. Required Accurate Phrases Check ---")
    for phrase in REQUIRED_PHRASES:
        found = phrase in full_text
        status = "[PASS]" if found else "[FAIL]"
        print(f"  {status} {phrase!r}")
        if not found:
            all_passed = False

    print("\n--- 2. Corrupted Artefacts Absence Check ---")
    for artefact in CORRUPTED_ARTEFACTS:
        found = artefact in full_text
        status = "[FAIL] (Still Present)" if found else "[PASS] (Eliminated)"
        print(f"  {status} {artefact!r}")
        if found:
            all_passed = False

    print("\n--- 3. Protected English Tokens Check ---")
    tor1_ok = "TOR-1" in full_text
    tor2_ok = "TOR-2" in full_text
    no_corrupt_tor = "ঞঙজ" not in full_text
    print(f"  {'[PASS]' if tor1_ok else '[FAIL]'} TOR-1 preserved")
    print(f"  {'[PASS]' if tor2_ok else '[FAIL]'} TOR-2 preserved")
    print(f"  {'[PASS]' if no_corrupt_tor else '[FAIL]'} No ঞঙজ corruption")
    if not (tor1_ok and tor2_ok and no_corrupt_tor):
        all_passed = False

    print("\n--- 4. Font Counts Check ---")
    font_counts = {}
    def tally_fonts(p):
        for r in p.runs:
            fn = r.font.name
            font_counts[fn] = font_counts.get(fn, 0) + 1

    for p in doc.paragraphs:
        tally_fonts(p)
    for t in doc.tables:
        def tally_table(tbl):
            for row in tbl.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        tally_fonts(p)
                    for nested in cell.tables:
                        tally_table(nested)
        tally_table(t)

    for fn, count in sorted(font_counts.items(), key=lambda x: str(x[0])):
        print(f"  Font: {fn!r:20s} | Runs: {count}")

    sutonny_remaining = font_counts.get("SutonnyMJ", 0)
    nikosh_present = font_counts.get("Nikosh", 0) > 0
    times_present = font_counts.get("Times New Roman", 0) > 0

    if sutonny_remaining > 0:
        print(f"  [FAIL] SutonnyMJ runs still remain: {sutonny_remaining}")
        all_passed = False
    else:
        print("  [PASS] Zero SutonnyMJ runs remaining")

    if not nikosh_present:
        print("  [FAIL] Nikosh runs missing")
        all_passed = False
    else:
        print("  [PASS] Nikosh runs active")

    if not times_present:
        print("  [FAIL] Times New Roman runs missing")
        all_passed = False
    else:
        print("  [PASS] Times New Roman runs preserved")

    print("\n" + "=" * 80)
    if all_passed:
        print("OVERALL RESULT: ALL TESTS PASSED SUCCESSFULLY")
    else:
        print("OVERALL RESULT: SOME TESTS FAILED")
    print("=" * 80)

    if not all_passed:
        sys.exit(1)

if __name__ == "__main__":
    main()

