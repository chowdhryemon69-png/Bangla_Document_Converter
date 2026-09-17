import unittest
from pathlib import Path

from docx import Document

from converter import normalize_latin_runs_in_document


class EnglishFontNormalizationTest(unittest.TestCase):

    def test_latin_runs_are_forced_to_times_new_roman(self):
        output = Path("output/test_english_font.docx")
        output.parent.mkdir(parents=True, exist_ok=True)

        doc = Document()
        p1 = doc.add_paragraph()
        r1 = p1.add_run("English Arial Text")
        r1.font.name = "Arial"
        r1.font.bold = True

        p2 = doc.add_paragraph()
        r2 = p2.add_run("English Calibri Text")
        r2.font.name = "Calibri"
        r2.font.italic = True

        p3 = doc.add_paragraph()
        r3 = p3.add_run("সংযুক্তি: TOR-1 এবং TOR-2")
        r3.font.name = "SutonnyMJ"

        p4 = doc.add_paragraph()
        r4 = p4.add_run("Existing TNR Text")
        r4.font.name = "Times New Roman"
        r4.font.bold = True

        doc.save(output)

        doc = Document(output)
        normalize_latin_runs_in_document(doc)
        doc.save(output)

        doc = Document(output)
        found_times = 0
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                text = run.text or ""
                if any(token in text for token in ["English", "TNR", "TOR-1", "TOR-2"]):
                    if run.font.name == "Times New Roman":
                        found_times += 1

        self.assertGreaterEqual(found_times, 4)


if __name__ == "__main__":
    unittest.main()
