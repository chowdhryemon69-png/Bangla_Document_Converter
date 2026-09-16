from docx import Document
document = Document()

document.add_heading("Bangla Document Converter", level=1)

document.add_paragraph(
    "This is our first test document."
)

document.add_paragraph(
    "এটি আমাদের প্রথম পরীক্ষা।"
)

document.save("test_document.docx")

print("SUCCESS: test_document.docx created.")