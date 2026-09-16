
import sys
from pathlib import Path
from docx import Document

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass


OUTPUT_FILE = Path(
    r"D:\Bangla_Document_Converter\output\Note_TEC_Unicode_Nikosh.docx"
)

SOURCE_FONT = "SutonnyMJ"
TARGET_FONT = "Nikosh"


def inspect_run(run, location, index):

    text = run.text or ""

    if not text:
        return

    print(
        f"{location} | Run {index:03d} | "
        f"Font={run.font.name!r} | "
        f"Size={run.font.size.pt if run.font.size else None} | "
        f"Text={text!r}"
    )


def inspect_paragraph(paragraph, location):

    for index, run in enumerate(
        paragraph.runs,
        start=1,
    ):
        inspect_run(
            run,
            location,
            index,
        )


def inspect_table(table, table_number):

    for row_number, row in enumerate(
        table.rows,
        start=1,
    ):

        for col_number, cell in enumerate(
            row.cells,
            start=1,
        ):

            location = (
                f"Table {table_number} "
                f"Cell [{row_number},{col_number}]"
            )

            print()
            print(location)
            print("-" * 70)

            print(
                f"Full text: {cell.text!r}"
            )

            for paragraph_number, paragraph in enumerate(
                cell.paragraphs,
                start=1,
            ):

                for run_number, run in enumerate(
                    paragraph.runs,
                    start=1,
                ):

                    inspect_run(
                        run,
                        f"{location} P{paragraph_number}",
                        run_number,
                    )

            for nested_number, nested_table in enumerate(
                cell.tables,
                start=1,
            ):

                inspect_table(
                    nested_table,
                    f"{table_number}.{nested_number}",
                )


def main():

    print("=" * 80)
    print("OUTPUT DOCUMENT VALIDATION")
    print("=" * 80)

    print(
        f"File: {OUTPUT_FILE}"
    )

    print()

    if not OUTPUT_FILE.exists():

        raise FileNotFoundError(
            f"Output file does not exist:\n{OUTPUT_FILE}"
        )

    document = Document(
        OUTPUT_FILE
    )

    print(
        f"Paragraphs: {len(document.paragraphs)}"
    )

    print(
        f"Tables    : {len(document.tables)}"
    )

    print()

    # ========================================================
    # Counters
    # ========================================================

    nikosh_runs = 0
    sutonny_runs = 0
    times_runs = 0
    other_runs = 0

    tor_ok = True

    tor_occurrences = []

    suspicious_legacy = []

    suspicious_wrong_tor = []

    # ========================================================
    # Inspect body paragraphs
    # ========================================================

    print("=" * 80)
    print("BODY PARAGRAPHS")
    print("=" * 80)

    for paragraph_number, paragraph in enumerate(
        document.paragraphs,
        start=1,
    ):

        if not paragraph.text:
            continue

        print()
        print("-" * 80)
        print(
            f"PARAGRAPH {paragraph_number}"
        )
        print("-" * 80)

        print(
            f"Full text: {paragraph.text!r}"
        )

        for run_number, run in enumerate(
            paragraph.runs,
            start=1,
        ):

            text = run.text or ""

            if run.font.name == TARGET_FONT:
                nikosh_runs += 1

            elif run.font.name == SOURCE_FONT:
                sutonny_runs += 1

            elif run.font.name == "Times New Roman":
                times_runs += 1

            elif text:
                other_runs += 1

            if "TOR-1" in text or "TOR-2" in text:
                tor_occurrences.append(
                    text
                )

            if "ঞঙজ" in text:
                suspicious_wrong_tor.append(
                    text
                )

            # Look for common signs that legacy Bijoy
            # text may still remain.
            legacy_markers = [
                "cÖ",
                "‡",
                "¨",
                "©",
                "ÿ",
                "we",
                "Kcv",
                "ZvwiL",
                "wbe©",
            ]

            if any(
                marker in text
                for marker in legacy_markers
            ):

                suspicious_legacy.append(
                    (
                        f"Paragraph {paragraph_number}",
                        text,
                    )
                )

            inspect_run(
                run,
                f"Paragraph {paragraph_number}",
                run_number,
            )

    # ========================================================
    # Inspect tables
    # ========================================================

    print()
    print("=" * 80)
    print("TABLES")
    print("=" * 80)

    for table_number, table in enumerate(
        document.tables,
        start=1,
    ):

        inspect_table(
            table,
            table_number,
        )

        # Count fonts inside tables.
        def count_table_fonts(tbl):

            nonlocal nikosh_runs
            nonlocal sutonny_runs
            nonlocal times_runs
            nonlocal other_runs
            nonlocal tor_ok

            for row in tbl.rows:

                for cell in row.cells:

                    for paragraph in cell.paragraphs:

                        for run in paragraph.runs:

                            text = run.text or ""

                            if run.font.name == TARGET_FONT:
                                nikosh_runs += 1

                            elif run.font.name == SOURCE_FONT:
                                sutonny_runs += 1

                            elif run.font.name == "Times New Roman":
                                times_runs += 1

                            elif text:
                                other_runs += 1

                            if "TOR-1" in text or "TOR-2" in text:
                                tor_occurrences.append(
                                    text
                                )

                            if "ঞঙজ" in text:
                                suspicious_wrong_tor.append(
                                    text
                                )

                    for nested in cell.tables:
                        count_table_fonts(nested)

        count_table_fonts(table)

    # ========================================================
    # Summary
    # ========================================================

    print()
    print("=" * 80)
    print("VALIDATION SUMMARY")
    print("=" * 80)

    print(
        f"Nikosh runs       : {nikosh_runs}"
    )

    print(
        f"SutonnyMJ runs    : {sutonny_runs}"
    )

    print(
        f"Times New Roman   : {times_runs}"
    )

    print(
        f"Other-font runs   : {other_runs}"
    )

    print()

    # ========================================================
    # TOR validation
    # ========================================================

    print(
        "Protected English tokens:"
    )

    if tor_occurrences:

        for item in tor_occurrences:
            print(
                f"  FOUND: {item!r}"
            )

    else:

        print(
            "  No TOR-1 / TOR-2 occurrence found."
        )

    print()

    if suspicious_wrong_tor:

        print(
            "[FAIL] Converted TOR token detected:"
        )

        for item in suspicious_wrong_tor:
            print(
                f"  {item!r}"
            )

        tor_ok = False

    else:

        print(
            "[PASS] No 'ঞঙজ' corruption detected."
        )

    # ========================================================
    # Legacy text validation
    # ========================================================

    print()

    if suspicious_legacy:

        print(
            "[WARNING] Possible legacy Bijoy text remains:"
        )

        # Show maximum 20 examples.
        for location, text in suspicious_legacy[:20]:

            print(
                f"  {location}: {text!r}"
            )

        if len(suspicious_legacy) > 20:

            print(
                f"  ... and "
                f"{len(suspicious_legacy) - 20} more."
            )

    else:

        print(
            "[PASS] No obvious legacy Bijoy markers detected."
        )

    # ========================================================
    # Final assessment
    # ========================================================

    print()
    print("=" * 80)
    print("FINAL CHECK")
    print("=" * 80)

    if sutonny_runs == 0:

        print(
            "[PASS] No SutonnyMJ runs detected."
        )

    else:

        print(
            "[WARNING] SutonnyMJ runs still remain."
        )

    if nikosh_runs > 0:

        print(
            "[PASS] Nikosh runs detected."
        )

    else:

        print(
            "[FAIL] No Nikosh runs detected."
        )

    if tor_ok:

        print(
            "[PASS] TOR protection appears successful."
        )

    else:

        print(
            "[FAIL] TOR protection failed."
        )

    print()
    print(
        "Validation complete."
    )
    print("=" * 80)


if __name__ == "__main__":

    try:
        main()

    except Exception as exc:

        print()
        print("[ERROR]")
        print(str(exc))

