from docx import Document
from pathlib import Path


# ============================================================
# CONFIGURATION
# ============================================================

INPUT_FILE = Path("samples/Note_TEC.docx")


# ============================================================
# CHECK FILE
# ============================================================

if not INPUT_FILE.exists():
    print()
    print("=" * 70)
    print("ERROR: FILE NOT FOUND")
    print("=" * 70)
    print(f"Expected file: {INPUT_FILE}")
    print(f"Full path:    {INPUT_FILE.resolve()}")
    print()
    print("Please make sure Note_TEC.docx is inside:")
    print("D:\\Bangla_Document_Converter\\samples")
    print()
    raise SystemExit(1)


# ============================================================
# LOAD DOCUMENT
# ============================================================

try:
    document = Document(INPUT_FILE)
except Exception as error:
    print()
    print("=" * 70)
    print("ERROR OPENING DOCUMENT")
    print("=" * 70)
    print(error)
    print()
    raise SystemExit(1)


# ============================================================
# HEADER
# ============================================================

print()
print("=" * 80)
print("BANGLA DOCUMENT CONVERTER")
print("DOCX INSPECTOR")
print("=" * 80)

print(f"Input file : {INPUT_FILE}")
print(f"Full path  : {INPUT_FILE.resolve()}")
print()

print("Document loaded successfully.")
print()


# ============================================================
# DOCUMENT INFORMATION
# ============================================================

print("=" * 80)
print("DOCUMENT INFORMATION")
print("=" * 80)

print(f"Paragraphs : {len(document.paragraphs)}")
print(f"Tables     : {len(document.tables)}")


# ============================================================
# PARAGRAPH INSPECTION
# ============================================================

print()
print("=" * 80)
print("PARAGRAPHS")
print("=" * 80)


non_empty_paragraphs = 0

for paragraph_number, paragraph in enumerate(
    document.paragraphs,
    start=1
):

    if not paragraph.text.strip():
        continue

    non_empty_paragraphs += 1

    print()
    print("-" * 80)
    print(f"PARAGRAPH {paragraph_number}")
    print("-" * 80)

    print("Full text:")
    print(repr(paragraph.text))

    print()
    print(f"Number of runs: {len(paragraph.runs)}")
    print()

    for run_number, run in enumerate(
        paragraph.runs,
        start=1
    ):

        font_name = run.font.name

        if run.font.size is not None:
            font_size = f"{run.font.size.pt} pt"
        else:
            font_size = "None"

        print(f"Run {run_number}:")
        print(f"  Font      : {font_name!r}")
        print(f"  Size      : {font_size}")
        print(f"  Bold      : {run.bold}")
        print(f"  Italic    : {run.italic}")
        print(f"  Underline : {run.underline}")
        print(f"  Text      : {run.text!r}")
        print()


# ============================================================
# TABLE INSPECTION
# ============================================================

print()
print("=" * 80)
print("TABLES")
print("=" * 80)

print(f"Number of tables: {len(document.tables)}")


for table_number, table in enumerate(
    document.tables,
    start=1
):

    print()
    print("-" * 80)
    print(f"TABLE {table_number}")
    print("-" * 80)

    print(f"Rows: {len(table.rows)}")

    for row_number, row in enumerate(
        table.rows,
        start=1
    ):

        for cell_number, cell in enumerate(
            row.cells,
            start=1
        ):

            if not cell.text.strip():
                continue

            print()
            print(
                f"Cell [{row_number},{cell_number}]"
            )

            print(
                f"  Text: {cell.text!r}"
            )

            for paragraph in cell.paragraphs:

                for run_number, run in enumerate(
                    paragraph.runs,
                    start=1
                ):

                    if not run.text:
                        continue

                    print(
                        f"  Run {run_number}: "
                        f"font={run.font.name!r}, "
                        f"bold={run.bold}, "
                        f"italic={run.italic}, "
                        f"text={run.text!r}"
                    )


# ============================================================
# SUMMARY
# ============================================================

print()
print("=" * 80)
print("INSPECTION COMPLETE")
print("=" * 80)

print(f"Non-empty paragraphs: {non_empty_paragraphs}")
print(f"Tables              : {len(document.tables)}")

print()
print("No changes were made to the original document.")
print()